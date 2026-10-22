"""
docx_generator.py
สร้างรายงานสรุปงาน (Task Report) เป็นไฟล์ Word (.docx)
"""
import logging
from collections import Counter
from datetime import date
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor

logger = logging.getLogger(__name__)


def generate_report(tasks, output_path, author="663380577-8"):
    """สร้างรายงาน Word ที่มีสรุปสถานะ, ตารางงานทั้งหมด และงานที่ยังไม่เสร็จ"""
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)

    doc = Document()

    title = doc.add_heading("Enhanced Task Manager - Project Report", level=0)
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER

    info = doc.add_paragraph()
    info.add_run("Author: ").bold = True
    info.add_run(author)
    info.add_run("\nGenerated on: ").bold = True
    info.add_run(date.today().isoformat())

    # 1. สรุปภาพรวม
    doc.add_heading("1. Summary", level=1)
    status_count = Counter(t["status"] for t in tasks)
    done = status_count.get("Done", 0)
    progress = done / len(tasks) * 100 if tasks else 0

    p = doc.add_paragraph()
    p.add_run(f"Total tasks: {len(tasks)}    ")
    run = p.add_run(f"Progress: {progress:.0f}% complete")
    run.bold = True
    run.font.color.rgb = RGBColor(0x2F, 0x54, 0x96)

    for status in ("Done", "In Progress", "To Do"):
        doc.add_paragraph(f"{status}: {status_count.get(status, 0)} task(s)", style="List Bullet")

    # 2. ตารางงานทั้งหมด
    doc.add_heading("2. All Tasks", level=1)
    headers = ["ID", "Task", "Category", "Priority", "Status", "Due Date"]
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Light Grid Accent 1"
    for cell, text in zip(table.rows[0].cells, headers):
        cell.text = text

    for t in tasks:
        cells = table.add_row().cells
        values = [t["id"], t["title"], t["category"], t["priority"], t["status"], t["due_date"]]
        for cell, value in zip(cells, values):
            cell.text = str(value)

    # 3. งานที่ยังไม่เสร็จ เรียงตามวันครบกำหนด
    doc.add_heading("3. Pending Tasks (by due date)", level=1)
    pending = sorted((t for t in tasks if t["status"] != "Done"), key=lambda t: t["due_date"])
    for t in pending:
        item = doc.add_paragraph(style="List Number")
        item.add_run(t["title"]).bold = True
        item.add_run(f" - {t['priority']} priority, due {t['due_date']} ({t['status']})")

    note = doc.add_paragraph("This report was generated automatically by the Enhanced Task Manager.")
    note.runs[0].italic = True
    note.runs[0].font.size = Pt(9)

    doc.save(output_path)
    logger.info("Word report saved: %s", output_path)
    return output_path
