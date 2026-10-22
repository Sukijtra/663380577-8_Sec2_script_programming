# agentic-spreadsheet-processor/src/spreadsheet_agent.py
import json
import datetime
from copy import copy

import openpyxl
from openpyxl.chart import BarChart, LineChart, PieChart, Reference
from openpyxl.formatting.rule import ColorScaleRule, DataBarRule, IconSetRule, FormulaRule
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter, range_boundaries
from openpyxl.worksheet.datavalidation import DataValidation

from .utils import setup_logging
from .google_sheets_utils import GoogleSheetsUtils

logger = setup_logging(__name__)

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill(start_color="2F5496", end_color="2F5496", fill_type="solid")


class SpreadsheetAgent:
    """
    An agent that performs a series of operations on an Excel workbook
    based on a provided configuration.
    """
    def __init__(self, config):
        self.config = config
        self.workbook = None
        self.audit_log = []
        self.start_time = datetime.datetime.now()

        # Task dispatcher: maps task "type" from the config to a handler method
        self.handlers = {
            "copy_data": self._process_copy_data_task,
            "calculate_column": self._process_calculate_column_task,
            "conditional_format": self._process_conditional_format_task,
            "chart": self._process_chart_task,
            "summary": self._process_summary_task,
            "data_validation": self._process_data_validation_task,
            "protect_sheet": self._process_protect_sheet_task,
        }
        logger.info("Spreadsheet Agent initialized.")

    def _log_audit(self, status, message, task_type="N/A", details=None):
        """Adds an entry to the audit log."""
        self.audit_log.append({
            "timestamp": datetime.datetime.now().isoformat(timespec="seconds"),
            "status": status,
            "task_type": task_type,
            "message": message,
            "details": details if details else {}
        })
        if status in ("ERROR", "CRITICAL"):
            logger.error(f"AUDIT LOG - {status}: {message} | Task: {task_type} | Details: {details}")
        else:
            logger.info(f"AUDIT LOG - {status}: {message} | Task: {task_type}")

    def _load_workbook(self):
        """Loads the input workbook."""
        input_file = self.config["input_file"]
        try:
            self.workbook = openpyxl.load_workbook(input_file)
            self._log_audit("SUCCESS", f"Workbook '{input_file}' loaded.", "load_workbook")
            return True
        except FileNotFoundError:
            self._log_audit("ERROR", f"Input file not found: {input_file}", "load_workbook")
            return False
        except Exception as e:
            self._log_audit("ERROR", f"Failed to load workbook '{input_file}': {e}", "load_workbook")
            return False

    def _get_sheet(self, sheet_name, create=True):
        """Safely gets a worksheet from the workbook, creating it if allowed."""
        if sheet_name in self.workbook.sheetnames:
            return self.workbook[sheet_name]
        if not create:
            return None
        self._log_audit("INFO", f"Sheet '{sheet_name}' not found. Creating new sheet.", "get_sheet")
        return self.workbook.create_sheet(sheet_name)

    def _style_header(self, sheet):
        """Makes row 1 bold with a colored background."""
        for cell in sheet[1]:
            if cell.value is not None:
                cell.font = HEADER_FONT
                cell.fill = HEADER_FILL

    def _auto_fit_columns(self, sheet, max_width=60):
        """Adjusts column widths to fit their content."""
        for col_cells in sheet.columns:
            length = max((len(str(c.value)) for c in col_cells if c.value is not None), default=0)
            sheet.column_dimensions[get_column_letter(col_cells[0].column)].width = min(length + 3, max_width)

    def _reference(self, default_sheet, cell_range):
        """Builds a chart Reference from "A1:B5" or a sheet-qualified range like "'Raw Data'!A1:B5"."""
        sheet = default_sheet
        if "!" in cell_range:
            sheet_name, cell_range = cell_range.rsplit("!", 1)
            sheet = self.workbook[sheet_name.strip("'")]
        min_col, min_row, max_col, max_row = range_boundaries(cell_range)
        return Reference(sheet, min_col=min_col, min_row=min_row, max_col=max_col, max_row=max_row)

    # --- Task handlers ---
    def _process_copy_data_task(self, task):
        """Copies data (values and basic styles) from a source sheet to a destination sheet."""
        src_sheet = self._get_sheet(task["source_sheet"], create=False)
        if src_sheet is None:
            self._log_audit("ERROR", f"Source sheet '{task['source_sheet']}' not found.", "copy_data", task)
            return
        dest_sheet = self._get_sheet(task["dest_sheet"])

        for row in src_sheet.iter_rows():
            for cell in row:
                new_cell = dest_sheet[cell.coordinate]
                new_cell.value = cell.value
                if cell.has_style:
                    new_cell.font = copy(cell.font)
                    new_cell.number_format = cell.number_format
        self._style_header(dest_sheet)
        self._log_audit("SUCCESS", f"Copied {src_sheet.max_row} rows from '{task['source_sheet']}' to '{task['dest_sheet']}'.", "copy_data")

    def _process_calculate_column_task(self, task):
        """Writes an Excel formula into every data row of a target column. {row} is the row number."""
        sheet = self._get_sheet(task["sheet"])
        column = task["target_column"]
        start_row = task.get("start_row", 2)  # Assume header in row 1
        last_row = sheet.max_row

        sheet[f"{column}1"].value = task.get("header", "Calculated Column")
        sheet[f"{column}1"].font = HEADER_FONT
        sheet[f"{column}1"].fill = HEADER_FILL

        for row_num in range(start_row, last_row + 1):
            cell = sheet[f"{column}{row_num}"]
            cell.value = task["formula"].replace("{row}", str(row_num))
            if task.get("number_format"):
                cell.number_format = task["number_format"]
        self._log_audit("SUCCESS", f"Formula column '{column}' ({task['formula']}) written to rows {start_row}-{last_row} in '{task['sheet']}'.", "calculate_column")

    def _process_conditional_format_task(self, task):
        """Applies conditional formatting to a range."""
        sheet = self._get_sheet(task["sheet"])
        data_range = task["range"]
        rule_type = task["rule_type"]

        if rule_type == "color_scale":
            rule = ColorScaleRule(start_type='min', start_color=task.get("min_color", "F8696B"),
                                  mid_type='percentile', mid_value=50, mid_color=task.get("mid_color", "FFEB84"),
                                  end_type='max', end_color=task.get("max_color", "63BE7B"))
        elif rule_type == "data_bar":
            rule = DataBarRule(start_type='min', end_type='max', color=task.get("color", "5B9BD5"))
        elif rule_type == "icon_set":
            rule = IconSetRule(task.get("icon_set", '3TrafficLights1'), 'percent', task.get("values", [0, 33, 67]))
        elif rule_type == "expression":
            # Conditional-format formulas are stored without the leading "="
            formula = task["formula"].lstrip("=")
            rule = FormulaRule(formula=[formula],
                               font=Font(color=task.get("font_color", "9C0006"), bold=task.get("bold", False)),
                               fill=PatternFill(start_color=task.get("fill_color", "FFC7CE"),
                                                end_color=task.get("fill_color", "FFC7CE"), fill_type="solid"))
        else:
            self._log_audit("ERROR", f"Unsupported conditional format rule_type: {rule_type}", "conditional_format", task)
            return

        sheet.conditional_formatting.add(data_range, rule)
        self._log_audit("SUCCESS", f"Conditional formatting '{rule_type}' applied to '{data_range}' in '{task['sheet']}'.", "conditional_format")

    def _process_chart_task(self, task):
        """Creates a bar, line or pie chart from a data range."""
        sheet = self._get_sheet(task["sheet"])
        chart_classes = {"bar": BarChart, "line": LineChart, "pie": PieChart}
        chart_class = chart_classes.get(task["chart_type"])
        if chart_class is None:
            self._log_audit("ERROR", f"Unsupported chart type: {task['chart_type']}", "chart", task)
            return

        chart = chart_class()
        chart.title = task.get("title", "Chart")
        chart.width = task.get("width", 18)
        chart.height = task.get("height", 9)

        chart.add_data(self._reference(sheet, task["data_range"]), titles_from_data=True)  # First row is the series name
        if task.get("category_range"):
            chart.set_categories(self._reference(sheet, task["category_range"]))

        if task["chart_type"] != "pie":
            chart.x_axis.title = task.get("x_axis_title", "")
            chart.y_axis.title = task.get("y_axis_title", "")
            # Newer Excel versions hide axes unless "delete" is explicitly False
            chart.x_axis.delete = False
            chart.y_axis.delete = False

        sheet.add_chart(chart, task.get("top_left_cell", "H2"))
        self._log_audit("SUCCESS", f"{task['chart_type'].title()} chart '{task.get('title', 'Chart')}' added to '{task['sheet']}'.", "chart")

    def _process_summary_task(self, task):
        """
        Builds a summary sheet of label/formula pairs that aggregate a column of the source sheet.
        Each metric: {"label": "...", "function": "SUM|AVERAGE|MAX|MIN|COUNT|COUNTIF", "column": "D", "criteria": "..."}
        """
        sheet = self._get_sheet(task["sheet"])
        src = task["source_sheet"]
        last_row = self._get_sheet(src, create=False).max_row

        sheet.append(["Metric", "Value"])
        for metric in task["metrics"]:
            col = metric["column"]
            rng = f"'{src}'!{col}2:{col}{last_row}"
            func = metric["function"].upper()
            if func == "COUNTIF":
                formula = f'=COUNTIF({rng},"{metric["criteria"]}")'
            else:
                formula = f"={func}({rng})"
            sheet.append([metric["label"], formula])
            if metric.get("number_format"):
                sheet.cell(row=sheet.max_row, column=2).number_format = metric["number_format"]

        self._style_header(sheet)
        self._auto_fit_columns(sheet)
        self._log_audit("SUCCESS", f"Summary sheet '{task['sheet']}' created with {len(task['metrics'])} metrics.", "summary")

    def _process_data_validation_task(self, task):
        """Adds a data validation rule (e.g. a dropdown list) to a range."""
        sheet = self._get_sheet(task["sheet"])
        if task["validation_type"] == "list":
            dv = DataValidation(type="list", formula1='"' + ",".join(task["options"]) + '"', allow_blank=True)
        else:  # whole, decimal, date, textLength
            dv = DataValidation(type=task["validation_type"], operator=task.get("operator", "between"),
                                formula1=str(task.get("min", 0)), formula2=str(task.get("max", 0)))
        dv.error = task.get("error_message", "Invalid value")
        dv.errorTitle = "Invalid input"
        dv.prompt = task.get("prompt", "")
        dv.showInputMessage = bool(task.get("prompt"))
        sheet.add_data_validation(dv)
        dv.add(task["range"])

        if task.get("header_cell"):
            sheet[task["header_cell"]].value = task.get("header", "Input")
            sheet[task["header_cell"]].font = HEADER_FONT
            sheet[task["header_cell"]].fill = HEADER_FILL
        self._log_audit("SUCCESS", f"Data validation '{task['validation_type']}' added to '{task['range']}' in '{task['sheet']}'.", "data_validation")

    def _process_protect_sheet_task(self, task):
        """Protects a sheet so formulas cannot be changed by accident."""
        sheet = self._get_sheet(task["sheet"])
        sheet.protection.sheet = True
        if task.get("password"):
            sheet.protection.password = task["password"]
        self._log_audit("SUCCESS", f"Sheet '{task['sheet']}' protected.", "protect_sheet")

    # --- Reporting ---
    def _add_audit_log_sheet(self):
        """Adds a sheet to the workbook containing the audit log."""
        audit_sheet = self._get_sheet("Audit Log", create=False) or self.workbook.create_sheet("Audit Log")
        audit_sheet.append(["Timestamp", "Status", "Task Type", "Message", "Details"])
        status_fills = {"SUCCESS": "C6EFCE", "WARNING": "FFEB9C", "ERROR": "FFC7CE", "CRITICAL": "FFC7CE"}

        for entry in self.audit_log:
            audit_sheet.append([entry["timestamp"], entry["status"], entry["task_type"],
                                entry["message"], json.dumps(entry["details"], ensure_ascii=False)])
            color = status_fills.get(entry["status"])
            if color:
                audit_sheet.cell(row=audit_sheet.max_row, column=2).fill = PatternFill(
                    start_color=color, end_color=color, fill_type="solid")

        self._style_header(audit_sheet)
        self._auto_fit_columns(audit_sheet)
        logger.info("Audit log sheet created.")

    def _upload_to_google_sheets(self):
        """Optionally uploads the processed sheets to Google Sheets."""
        gs_config = self.config.get("google_sheets", {})
        if not gs_config.get("enabled"):
            self._log_audit("INFO", "Google Sheets upload disabled in config. Skipped.", "google_sheets")
            return
        gs = GoogleSheetsUtils(gs_config.get("credentials_file", "configs/service_account.json"))
        if not gs.connect():
            self._log_audit("WARNING", "Could not connect to Google Sheets. Upload skipped.", "google_sheets")
            return
        try:
            for sheet_name in gs_config.get("sheets", []):
                gs.upload_worksheet(gs_config["spreadsheet_id"], self.workbook[sheet_name])
            self._log_audit("SUCCESS", f"Uploaded {gs_config.get('sheets')} to Google Sheets.", "google_sheets")
        except Exception as e:
            self._log_audit("ERROR", f"Google Sheets upload failed: {e}", "google_sheets")

    def run(self):
        """Executes the defined tasks from the configuration."""
        if not self._load_workbook():
            self._log_audit("CRITICAL", "Agent could not start due to workbook loading failure.", "run")
            return False

        for i, task in enumerate(self.config["tasks"], start=1):
            task_type = task.get("type")
            logger.info(f"Executing Task {i}: Type='{task_type}' Sheet='{task.get('sheet', task.get('dest_sheet'))}'")
            handler = self.handlers.get(task_type)
            if handler is None:
                self._log_audit("WARNING", f"Unsupported task type encountered: {task_type}", "task_execution", task)
                continue
            try:
                handler(task)
            except Exception as e:
                # One failing task should not stop the rest of the workflow
                self._log_audit("ERROR", f"Task {i} ({task_type}) failed: {e}", task_type, task)

        for sheet_name in self.config.get("auto_fit_sheets", []):
            if sheet_name in self.workbook.sheetnames:
                self._auto_fit_columns(self.workbook[sheet_name])

        self._upload_to_google_sheets()

        duration = (datetime.datetime.now() - self.start_time).total_seconds()
        self._log_audit("INFO", f"Agent execution completed in {duration:.2f} seconds.", "run_summary")

        # Add final audit log sheet and save the workbook
        self._add_audit_log_sheet()
        output_file = self.config["output_file"]
        try:
            self.workbook.save(output_file)
            logger.info(f"Final workbook saved to: {output_file}")
            return True
        except Exception as e:
            logger.critical(f"Error saving output workbook '{output_file}': {e}")
            return False
